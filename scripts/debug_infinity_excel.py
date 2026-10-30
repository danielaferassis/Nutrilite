from pathlib import Path

from openpyxl import load_workbook

path = Path(r"c:\Users\User1\Documents\GitHub\Nutrilite\ii9A18.tmp.xlsx")

wb = load_workbook(path, read_only=True, data_only=True)
print('SHEETS:', wb.sheetnames)

for nome in wb.sheetnames:
    ws = wb[nome]
    print(f'\nSHEET: {nome} :: rows={ws.max_row} cols={ws.max_column}')
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 10), values_only=True):
        print(row)
